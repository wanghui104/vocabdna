#!/usr/bin/env python
"""Update TOEFL 5400 ingest progress after a verified batch."""

from __future__ import annotations

import argparse
import json
import re
import sys
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path
from typing import Any

import openpyxl

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

DEFAULT_PROGRESS = Path("data/vocab/v1/toefl-5400-ingest-progress.json")
REQUIRED_HEADERS = ("No.", "Word", "POS", "Meaning", "Source_Page")


@dataclass
class LogicalWord:
    sheet: str
    start_row: int
    end_row: int
    number: int | None
    term: str
    rows: list[int] = field(default_factory=list)


def main() -> int:
    args = parse_args()
    root = Path(args.project_root).resolve()
    progress_path = resolve_project_path(root, Path(args.progress))
    progress = read_json(progress_path)

    sheet = args.sheet or progress.get("nextStartSheet")
    start_row = args.start_row or int(progress.get("nextStartRow", 0))
    if not args.end_row:
        raise SystemExit("--end-row is required so progress is advanced only to an explicit verified boundary.")
    end_row = args.end_row

    workbook_path = resolve_project_path(root, Path(args.workbook or progress["workbook"]))
    wb = openpyxl.load_workbook(workbook_path, data_only=True, read_only=True)
    sheets = ordered_range_sheets(wb.sheetnames)
    if sheet not in sheets:
        raise SystemExit(f"Sheet not found: {sheet}")

    words_by_sheet = {
        name: parse_sheet(wb[name], name, max_row=(end_row + 100 if name == sheet else 150))
        for name in sheets[sheets.index(sheet) : sheets.index(sheet) + 2]
    }
    batch_words = [w for w in words_by_sheet[sheet] if w.end_row >= start_row and w.start_row <= end_row]
    if not batch_words:
        raise SystemExit(f"No workbook words found in {sheet} rows {start_row}-{end_row}")

    next_start = find_next_start(words_by_sheet, sheets, sheet, end_row)
    proposed = dict(progress)
    proposed.update(
        {
            "completedThroughSheet": sheet,
            "completedThroughRow": end_row,
            "completedThroughNo": batch_words[-1].number,
            "nextStartSheet": next_start["sheet"],
            "nextStartRow": next_start["row"],
            "nextStartNo": next_start["no"],
            "lastBatch": {
                "sheet": sheet,
                "startRow": start_row,
                "endRow": end_row,
                "excelRowsCovered": count_rows(batch_words),
                "logicalWordsProcessed": len(batch_words),
                "firstWord": batch_words[0].term,
                "lastWord": batch_words[-1].term,
                "commit": args.commit,
            },
            "updatedAt": args.updated_at or date.today().isoformat(),
        }
    )

    if args.write:
        progress_path.write_text(json.dumps(proposed, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    result = {"dryRun": not args.write, "progress": proposed}
    if args.json:
        print(json.dumps(result, ensure_ascii=False, indent=2))
    else:
        print("TOEFL 5400 progress update" + (" (dry run)" if not args.write else ""))
        print(f"Completed: {sheet} row {end_row} / No. {batch_words[-1].number} / {batch_words[-1].term}")
        print(f"Rows covered: {count_rows(batch_words)}; logical words: {len(batch_words)}")
        print(f"Next start: {next_start['sheet']} row {next_start['row']} / No. {next_start['no']}")
    return 0


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--project-root", default=".")
    parser.add_argument("--progress", default=str(DEFAULT_PROGRESS))
    parser.add_argument("--workbook")
    parser.add_argument("--sheet")
    parser.add_argument("--start-row", type=int)
    parser.add_argument("--end-row", type=int)
    parser.add_argument("--commit")
    parser.add_argument("--updated-at")
    parser.add_argument("--write", action="store_true", help="Write the progress file. Default is dry-run.")
    parser.add_argument("--json", action="store_true")
    return parser.parse_args()


def read_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def resolve_project_path(project_root: Path, path: Path) -> Path:
    return path if path.is_absolute() else project_root / path


def ordered_range_sheets(sheet_names: list[str]) -> list[str]:
    def key(name: str) -> tuple[int, str]:
        match = re.match(r"^(\d+)-\d+$", name)
        return (int(match.group(1)) if match else 10**9, name)

    return sorted(sheet_names, key=key)


def parse_sheet(worksheet: Any, sheet_name: str, max_row: int | None = None) -> list[LogicalWord]:
    headers = [cell.value for cell in worksheet[1]]
    missing = [h for h in REQUIRED_HEADERS if h not in headers]
    if missing:
        raise ValueError(f"{sheet_name} missing headers: {', '.join(missing)}")
    index = {h: headers.index(h) + 1 for h in REQUIRED_HEADERS}
    logical_words: list[LogicalWord] = []
    current: LogicalWord | None = None
    for row_number in range(2, (max_row or worksheet.max_row) + 1):
        raw = {h: worksheet.cell(row_number, index[h]).value for h in REQUIRED_HEADERS}
        if all(v is None for v in raw.values()):
            continue
        term = clean(raw["Word"])
        if term:
            current = LogicalWord(sheet_name, row_number, row_number, to_int(raw["No."]), term)
            logical_words.append(current)
        elif current is None:
            raise ValueError(f"{sheet_name} row {row_number} is a continuation row without a word")
        assert current is not None
        current.end_row = row_number
        current.rows.append(row_number)
    return logical_words


def to_int(value: Any) -> int | None:
    return None if value is None or value == "" else int(value)


def clean(value: Any) -> str:
    return "" if value is None else str(value).strip()


def find_next_start(words_by_sheet: dict[str, list[LogicalWord]], sheets: list[str], end_sheet: str, end_row: int) -> dict[str, int | str | None]:
    for sheet in sheets[sheets.index(end_sheet) :]:
        for word in words_by_sheet.get(sheet, []):
            if sheet == end_sheet and word.start_row <= end_row:
                continue
            return {"sheet": sheet, "row": word.start_row, "no": word.number}
    return {"sheet": None, "row": None, "no": None}


def count_rows(words: list[LogicalWord]) -> int:
    return sum(len(w.rows) for w in words)


if __name__ == "__main__":
    raise SystemExit(main())
