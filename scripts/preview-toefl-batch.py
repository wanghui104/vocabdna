#!/usr/bin/env python
"""Preview the next TOEFL 5400 workbook batch without changing project files."""

from __future__ import annotations

import argparse
import json
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import openpyxl


if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

DEFAULT_PROGRESS = Path("data/vocab/v1/toefl-5400-ingest-progress.json")
DEFAULT_WORDS = Path("data/vocab/v1/words.seed-50.json")
REQUIRED_HEADERS = ("No.", "Word", "POS", "Meaning", "Source_Page")


@dataclass
class SourceRow:
    sheet: str
    row: int
    number: int | None
    term: str
    pos: str
    meaning: str
    source_page: str


@dataclass
class LogicalWord:
    sheet: str
    start_row: int
    end_row: int
    number: int | None
    term: str
    rows: list[SourceRow] = field(default_factory=list)

    @property
    def slug(self) -> str:
        return slugify(self.term)

    @property
    def pos(self) -> list[str]:
        return unique(row.pos for row in self.rows if row.pos)

    @property
    def meanings(self) -> list[str]:
        return [row.meaning for row in self.rows if row.meaning]


@dataclass
class Cluster:
    words: list[LogicalWord]

    @property
    def sheet(self) -> str:
        return self.words[0].sheet

    @property
    def start_row(self) -> int:
        return self.words[0].start_row

    @property
    def end_row(self) -> int:
        return self.words[-1].end_row

    @property
    def terms(self) -> list[str]:
        return [word.term for word in self.words]


def main() -> int:
    args = parse_args()
    project_root = Path(args.project_root).resolve()
    progress_path = resolve_project_path(project_root, Path(args.progress))
    progress = read_json(progress_path)

    workbook_path = resolve_project_path(project_root, Path(args.workbook or progress["workbook"]))
    target_rows = args.target_rows or int(progress.get("batchTargetRows", 50))
    start_sheet = args.sheet or progress["nextStartSheet"]
    start_row = args.row or int(progress["nextStartRow"])

    workbook = openpyxl.load_workbook(workbook_path, data_only=True, read_only=True)
    sheet_names = ordered_range_sheets(workbook.sheetnames)
    if start_sheet not in sheet_names:
        raise SystemExit(f"Sheet not found: {start_sheet}")

    sheet_window = sheet_names[sheet_names.index(start_sheet) : sheet_names.index(start_sheet) + 2]
    words_by_sheet = {sheet_name: parse_sheet(workbook[sheet_name], sheet_name, max_row=start_row + target_rows + 100) for sheet_name in sheet_window}
    clusters_by_sheet = {sheet_name: cluster_words(words) for sheet_name, words in words_by_sheet.items()}

    batch, warnings = select_batch(clusters_by_sheet, sheet_names, start_sheet, start_row, target_rows)
    logical_words = [word for cluster in batch for word in cluster.words]
    end_word = logical_words[-1]
    next_start = find_next_start(words_by_sheet, sheet_names, end_word.sheet, end_word.end_row)
    existing_slugs = load_existing_slugs(project_root / DEFAULT_WORDS)

    result = {
        "workbook": str(workbook_path.relative_to(project_root)),
        "targetRows": target_rows,
        "start": {"sheet": start_sheet, "row": start_row},
        "end": {"sheet": end_word.sheet, "row": end_word.end_row, "no": end_word.number, "word": end_word.term},
        "nextStart": next_start,
        "excelRowsCovered": count_rows(logical_words),
        "logicalWords": len(logical_words),
        "clusters": [
            {"sheet": cluster.sheet, "startRow": cluster.start_row, "endRow": cluster.end_row, "terms": cluster.terms}
            for cluster in batch
        ],
        "words": [
            {
                "sheet": word.sheet,
                "startRow": word.start_row,
                "endRow": word.end_row,
                "no": word.number,
                "term": word.term,
                "slug": word.slug,
                "pos": word.pos,
                "meanings": word.meanings,
                "alreadyInSource": word.slug in existing_slugs,
            }
            for word in logical_words
        ],
        "warnings": warnings,
    }

    if args.json:
        print(json.dumps(result, ensure_ascii=False, indent=2))
    else:
        print_human_preview(result)
    return 0


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--project-root", default=".", help="Project root. Defaults to current directory.")
    parser.add_argument("--progress", default=str(DEFAULT_PROGRESS), help="Progress JSON path.")
    parser.add_argument("--workbook", help="Workbook path. Defaults to progress.workbook.")
    parser.add_argument("--sheet", help="Override start sheet. Defaults to progress.nextStartSheet.")
    parser.add_argument("--row", type=int, help="Override start Excel row. Defaults to progress.nextStartRow.")
    parser.add_argument("--target-rows", type=int, help="Override target source row count.")
    parser.add_argument("--json", action="store_true", help="Print machine-readable JSON.")
    return parser.parse_args()


def read_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def resolve_project_path(project_root: Path, path: Path) -> Path:
    return path if path.is_absolute() else project_root / path


def ordered_range_sheets(sheet_names: list[str]) -> list[str]:
    def key(name: str) -> tuple[int, str]:
        match = re.match(r"^(\d+)-\d+$", name)
        return (int(match.group(1)) if match else 10**9, name)

    return sorted(sheet_names, key=key)


def parse_sheet(worksheet: Any, sheet_name: str, max_row: int | None = None) -> list[LogicalWord]:
    headers = [cell.value for cell in worksheet[1]]
    missing = [header for header in REQUIRED_HEADERS if header not in headers]
    if missing:
        raise ValueError(f"{sheet_name} missing headers: {', '.join(missing)}")

    index = {header: headers.index(header) + 1 for header in REQUIRED_HEADERS}
    logical_words: list[LogicalWord] = []
    current: LogicalWord | None = None

    last_row = max_row if max_row is not None else worksheet.max_row
    for row_number in range(2, last_row + 1):
        raw_no = worksheet.cell(row_number, index["No."]).value
        raw_word = worksheet.cell(row_number, index["Word"]).value
        raw_pos = worksheet.cell(row_number, index["POS"]).value
        raw_meaning = worksheet.cell(row_number, index["Meaning"]).value
        raw_source = worksheet.cell(row_number, index["Source_Page"]).value

        if all(value is None for value in (raw_no, raw_word, raw_pos, raw_meaning, raw_source)):
            continue

        term = clean(raw_word)
        if term:
            current = LogicalWord(sheet_name, row_number, row_number, to_int(raw_no), term)
            logical_words.append(current)
        elif current is None:
            raise ValueError(f"{sheet_name} row {row_number} is a continuation row without a word")

        assert current is not None
        current.end_row = row_number
        current.rows.append(
            SourceRow(
                sheet_name,
                row_number,
                to_int(raw_no) if raw_no is not None else current.number,
                current.term,
                clean(raw_pos),
                clean(raw_meaning),
                clean(raw_source),
            )
        )

    return logical_words


def to_int(value: Any) -> int | None:
    if value is None or value == "":
        return None
    return int(value)


def clean(value: Any) -> str:
    return "" if value is None else str(value).strip()


def slugify(term: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", term.lower()).strip("-")


def unique(values: Any) -> list[str]:
    seen: set[str] = set()
    result: list[str] = []
    for value in values:
        if value not in seen:
            seen.add(value)
            result.append(value)
    return result


def cluster_words(words: list[LogicalWord]) -> list[Cluster]:
    if not words:
        return []
    clusters: list[Cluster] = []
    current = Cluster([words[0]])
    for word in words[1:]:
        if are_obvious_family(current.words[-1].term, word.term):
            current.words.append(word)
        else:
            clusters.append(current)
            current = Cluster([word])
    clusters.append(current)
    return clusters


def are_obvious_family(left: str, right: str) -> bool:
    left_keys = family_keys(left)
    right_keys = family_keys(right)
    if left_keys & right_keys:
        return True

    left_compact = compact(left)
    right_compact = compact(right)
    if left_compact.startswith("un") and left_compact[2:] == right_compact:
        return True
    if right_compact.startswith("un") and right_compact[2:] == left_compact:
        return True

    return False


def family_keys(term: str) -> set[str]:
    value = compact(term)
    values = {value}
    if "addition" in value:
        values.update({"add", "addit", "addition"})
    if value.startswith("inadditionto"):
        values.update({"add", "addit", "addition"})
    if value.startswith("adventur"):
        values.update({"adventure", "adventur"})
    if value.startswith("administr"):
        values.update({"administer", "administr"})
    if value.startswith("altern"):
        values.add("altern")
    if value.startswith("ambit"):
        values.add("ambit")
    if value.startswith("amaz"):
        values.add("amaz")

    for suffix in (
        "able",
        "ible",
        "ation",
        "ition",
        "tion",
        "sion",
        "itive",
        "ative",
        "ive",
        "ial",
        "al",
        "ment",
        "ed",
        "ing",
        "er",
        "or",
        "ous",
        "ly",
        "ity",
    ):
        if value.endswith(suffix) and len(value) > len(suffix) + 2:
            values.add(value[: -len(suffix)])
    if value.endswith("ation") and len(value) > 7:
        values.add(value[:-5] + "e")
    if value.endswith("ive") and len(value) > 5:
        values.add(value[:-3])
    return {item for item in values if len(item) >= 3}


def compact(term: str) -> str:
    return re.sub(r"[^a-z0-9]", "", term.lower())


def common_prefix_len(left: str, right: str) -> int:
    count = 0
    for left_char, right_char in zip(left, right):
        if left_char != right_char:
            break
        count += 1
    return count


def select_batch(
    clusters_by_sheet: dict[str, list[Cluster]],
    sheet_names: list[str],
    start_sheet: str,
    start_row: int,
    target_rows: int,
) -> tuple[list[Cluster], list[str]]:
    selected: list[Cluster] = []
    warnings: list[str] = []
    started = False

    for sheet_name in sheet_names[sheet_names.index(start_sheet) :]:
        for cluster in clusters_by_sheet[sheet_name]:
            if not started:
                if cluster.end_row < start_row:
                    continue
                started = True
                if cluster.start_row < start_row:
                    warnings.append(
                        f"Start row {start_row} falls inside '{cluster.words[0].term}' "
                        f"({sheet_name} rows {cluster.start_row}-{cluster.end_row}); including the full word."
                    )
            selected.append(cluster)
            if count_rows([word for item in selected for word in item.words]) >= target_rows:
                return selected, warnings

    if not selected:
        raise ValueError(f"No data rows found at or after {start_sheet} row {start_row}")
    warnings.append("Reached the end of the workbook before targetRows was met.")
    return selected, warnings


def count_rows(words: list[LogicalWord]) -> int:
    return sum(len(word.rows) for word in words)


def find_next_start(
    words_by_sheet: dict[str, list[LogicalWord]],
    sheet_names: list[str],
    end_sheet: str,
    end_row: int,
) -> dict[str, int | str | None]:
    for sheet_name in sheet_names[sheet_names.index(end_sheet) :]:
        for word in words_by_sheet.get(sheet_name, []):
            if sheet_name == end_sheet and word.start_row <= end_row:
                continue
            return {"sheet": sheet_name, "row": word.start_row, "no": word.number}
    return {"sheet": None, "row": None, "no": None}


def load_existing_slugs(words_path: Path) -> set[str]:
    if not words_path.exists():
        return set()
    doc = read_json(words_path)
    return {word["slug"] for word in doc.get("words", []) if "slug" in word}


def print_human_preview(result: dict[str, Any]) -> None:
    end = result["end"]
    next_start = result["nextStart"]
    print("TOEFL 5400 batch preview")
    print(f"Workbook: {result['workbook']}")
    print(f"Start: {result['start']['sheet']} row {result['start']['row']}")
    print(f"End: {end['sheet']} row {end['row']} / No. {end['no']} / {end['word']}")
    print(f"Rows covered: {result['excelRowsCovered']}; logical words: {result['logicalWords']}")
    print(f"Next start: {next_start['sheet']} row {next_start['row']} / No. {next_start['no']}")
    print("")
    print("Clusters:")
    for cluster in result["clusters"]:
        print(f"- {cluster['sheet']} rows {cluster['startRow']}-{cluster['endRow']}: {', '.join(cluster['terms'])}")
    print("")
    print("Words:")
    for word in result["words"]:
        exists = "existing" if word["alreadyInSource"] else "missing"
        meanings = " | ".join(word["meanings"])
        print(
            f"- {word['sheet']} row {word['startRow']}-{word['endRow']} "
            f"No. {word['no']} {word['term']} [{exists}]: {meanings}"
        )
    if result["warnings"]:
        print("")
        print("Warnings:")
        for warning in result["warnings"]:
            print(f"- {warning}")


if __name__ == "__main__":
    raise SystemExit(main())





