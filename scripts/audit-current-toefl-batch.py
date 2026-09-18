#!/usr/bin/env python
"""Audit the latest TOEFL 5400 batch for relationship and generated-file consistency."""

from __future__ import annotations

import argparse
import json
import re
import sys
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import openpyxl

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

DEFAULT_PROGRESS = Path("data/vocab/v1/toefl-5400-ingest-progress.json")
WORDS_SOURCE = Path("data/vocab/v1/words.seed-50.json")
ROOTS_SOURCE = Path("data/vocab/v1/roots.seed.json")
WORDS_OUTPUT = Path("data/words")
MORPHEMES_OUTPUT = Path("data/morphemes")
REQUIRED_HEADERS = ("No.", "Word", "POS", "Meaning", "Source_Page")
RELATION_FIELDS = ("derivatives", "sameFamily", "confusables")


@dataclass
class LogicalWord:
    sheet: str
    start_row: int
    end_row: int
    number: int | None
    term: str
    rows: list[int] = field(default_factory=list)

    @property
    def slug(self) -> str:
        return slugify(self.term)


def main() -> int:
    args = parse_args()
    root = Path(args.project_root).resolve()
    progress = read_json(resolve_project_path(root, Path(args.progress)))
    last_batch = progress.get("lastBatch") or {}
    sheet = args.sheet or last_batch.get("sheet") or progress.get("completedThroughSheet")
    start_row = args.start_row or int(last_batch.get("startRow") or progress.get("nextStartRow", 0))
    end_row = args.end_row or int(last_batch.get("endRow") or progress.get("completedThroughRow", 0))

    workbook_path = resolve_project_path(root, Path(args.workbook or progress["workbook"]))
    wb = openpyxl.load_workbook(workbook_path, data_only=True, read_only=True)
    words = parse_sheet(wb[sheet], sheet, max_row=end_row + 5)
    batch_words = [w for w in words if w.end_row >= start_row and w.start_row <= end_row]

    words_doc = read_json(root / WORDS_SOURCE)
    roots_doc = read_json(root / ROOTS_SOURCE)
    entries = words_doc.get("words", [])
    roots = roots_doc.get("roots", [])
    by_slug = {w["slug"]: w for w in entries}
    root_ids = {r["id"] for r in roots}
    source_slugs = [w.slug for w in batch_words]
    unique_slugs = sorted(set(source_slugs), key=source_slugs.index)

    issues: list[str] = []
    warnings: list[str] = []
    notes: list[str] = []

    duplicate_source = [slug for slug, count in Counter(source_slugs).items() if count > 1]
    for slug in duplicate_source:
        notes.append(f"Workbook repeats slug '{slug}' in this batch; expecting one canonical entry.")

    duplicate_canonical = [slug for slug, count in Counter(w.get("slug") for w in entries).items() if slug and count > 1]
    for slug in duplicate_canonical:
        issues.append(f"Duplicate canonical word slug: {slug}")

    for slug in unique_slugs:
        entry = by_slug.get(slug)
        if not entry:
            issues.append(f"Missing canonical word for batch slug: {slug}")
            continue
        if not (root / WORDS_OUTPUT / f"{slug}.yaml").exists():
            issues.append(f"Missing generated word YAML: data/words/{slug}.yaml")
        audit_roots(slug, entry, root_ids, root, issues, warnings)
        audit_relations(slug, entry, by_slug, issues, warnings)

    result = {
        "batch": {"sheet": sheet, "startRow": start_row, "endRow": end_row, "sourceRows": count_rows(batch_words), "logicalWords": len(batch_words)},
        "checkedSlugs": source_slugs,
        "uniqueCheckedSlugs": unique_slugs,
        "issues": issues,
        "warnings": warnings,
        "notes": notes,
    }
    if args.json:
        print(json.dumps(result, ensure_ascii=False, indent=2))
    else:
        print_human(result)
    return 1 if issues else 0


def audit_roots(slug: str, entry: dict[str, Any], root_ids: set[str], root: Path, issues: list[str], warnings: list[str]) -> None:
    entry_root_ids = entry.get("rootIds", [])
    if not entry_root_ids:
        issues.append(f"{slug} has no rootIds")
    for root_id in entry_root_ids:
        if root_id not in root_ids:
            issues.append(f"{slug} references missing rootId {root_id}")
        elif not (root / MORPHEMES_OUTPUT / f"{safe_yaml_basename(root_id.removeprefix('root:'))}.yaml").exists():
            warnings.append(f"{slug} rootId {root_id} has no generated morpheme YAML")
    for morpheme in entry.get("morphemes", []):
        root_id = morpheme.get("rootId")
        if root_id not in root_ids:
            issues.append(f"{slug}.{morpheme.get('form', 'morpheme')} references missing rootId {root_id}")
        if root_id and root_id not in entry_root_ids:
            issues.append(f"{slug}.{morpheme.get('form', 'morpheme')} rootId {root_id} is absent from rootIds")


def audit_relations(slug: str, entry: dict[str, Any], by_slug: dict[str, dict[str, Any]], issues: list[str], warnings: list[str]) -> None:
    relations = entry.get("relations", {})
    for field in RELATION_FIELDS:
        values = relations.get(field, [])
        duplicates = [item for item, count in Counter(values).items() if count > 1]
        for item in duplicates:
            issues.append(f"{slug}.relations.{field} repeats {item}")
        for target in values:
            target_entry = by_slug.get(target)
            if not target_entry:
                if field == "derivatives":
                    issues.append(f"{slug}.relations.derivatives points outside canonical data: {target}")
                else:
                    warnings.append(f"{slug}.relations.{field} points outside canonical data: {target}")
                continue
            if field == "derivatives" and slug not in target_entry.get("relations", {}).get("derivatives", []):
                issues.append(f"{slug} derivative {target} is not reciprocal")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--project-root", default=".")
    parser.add_argument("--progress", default=str(DEFAULT_PROGRESS))
    parser.add_argument("--workbook")
    parser.add_argument("--sheet")
    parser.add_argument("--start-row", type=int)
    parser.add_argument("--end-row", type=int)
    parser.add_argument("--json", action="store_true")
    return parser.parse_args()


def read_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def resolve_project_path(project_root: Path, path: Path) -> Path:
    return path if path.is_absolute() else project_root / path


def parse_sheet(worksheet: Any, sheet_name: str, max_row: int | None = None) -> list[LogicalWord]:
    headers = [cell.value for cell in worksheet[1]]
    missing = [h for h in REQUIRED_HEADERS if h not in headers]
    if missing:
        raise ValueError(f"{sheet_name} missing headers: {', '.join(missing)}")
    index = {h: headers.index(h) + 1 for h in REQUIRED_HEADERS}
    logical_words: list[LogicalWord] = []
    current: LogicalWord | None = None
    for row_number in range(2, (max_row or worksheet.max_row) + 1):
        raw_no = worksheet.cell(row_number, index["No."]).value
        raw_word = worksheet.cell(row_number, index["Word"]).value
        raw_pos = worksheet.cell(row_number, index["POS"]).value
        raw_meaning = worksheet.cell(row_number, index["Meaning"]).value
        raw_source = worksheet.cell(row_number, index["Source_Page"]).value
        if all(value is None for value in (raw_no, raw_word, raw_pos, raw_meaning, raw_source)):
            continue
        term = clean(raw_word)
        if term:
            current = LogicalWord(sheet_name, row_number, row_number, to_int(raw_no), term)
            logical_words.append(current)
        elif current is None:
            raise ValueError(f"{sheet_name} row {row_number} is a continuation row without a word")
        assert current is not None
        current.end_row = row_number
        current.rows.append(row_number)
    return logical_words


def to_int(value: Any) -> int | None:
    return None if value is None or value == "" else int(value)


def clean(value: Any) -> str:
    return "" if value is None else str(value).strip()


def slugify(term: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", term.lower()).strip("-")


def count_rows(words: list[LogicalWord]) -> int:
    return sum(len(w.rows) for w in words)



def safe_yaml_basename(root_id: str) -> str:
    reserved_windows_names = {
        "con",
        "prn",
        "aux",
        "nul",
        "com1",
        "com2",
        "com3",
        "com4",
        "com5",
        "com6",
        "com7",
        "com8",
        "com9",
        "lpt1",
        "lpt2",
        "lpt3",
        "lpt4",
        "lpt5",
        "lpt6",
        "lpt7",
        "lpt8",
        "lpt9",
    }
    return f"{root_id}-root" if root_id.lower() in reserved_windows_names else root_id

def print_human(result: dict[str, Any]) -> None:
    batch = result["batch"]
    print("TOEFL 5400 batch relationship audit")
    print(f"Batch: {batch['sheet']} rows {batch['startRow']}-{batch['endRow']} ({batch['sourceRows']} source rows, {batch['logicalWords']} logical words)")
    print(f"Unique canonical slugs checked: {len(result['uniqueCheckedSlugs'])}")
    for label in ("issues", "warnings", "notes"):
        if result[label]:
            print(f"\n{label.title()}:")
            for item in result[label]:
                print(f"- {item}")
    if not result["issues"]:
        print("\nOK: current batch relationships and generated files passed.")


if __name__ == "__main__":
    raise SystemExit(main())
