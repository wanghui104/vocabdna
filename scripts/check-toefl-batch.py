#!/usr/bin/env python
"""Check the most recent TOEFL 5400 batch against source and generated data."""

from __future__ import annotations

import argparse
import json
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import openpyxl


if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

DEFAULT_PROGRESS = Path("data/vocab/v1/toefl-5400-ingest-progress.json")
WORDS_SOURCE = Path("data/vocab/v1/words.seed-50.json")
ROOTS_SOURCE = Path("data/vocab/v1/roots.seed.json")
WORDS_OUTPUT = Path("data/words")
MORPHEMES_OUTPUT = Path("data/morphemes")
REQUIRED_HEADERS = ("No.", "Word", "POS", "Meaning", "Source_Page")


@dataclass
class LogicalWord:
    sheet: str
    start_row: int
    end_row: int
    number: int | None
    term: str
    rows: list[int] = field(default_factory=list)

    @property
    def slug(self) -> str:
        return slugify(self.term)


def main() -> int:
    args = parse_args()
    project_root = Path(args.project_root).resolve()
    progress_path = resolve_project_path(project_root, Path(args.progress))
    progress = read_json(progress_path)
    last_batch = progress.get("lastBatch") or {}

    workbook_path = resolve_project_path(project_root, Path(args.workbook or progress["workbook"]))
    sheet = args.sheet or last_batch.get("sheet") or progress.get("completedThroughSheet")
    start_row = args.start_row or int(last_batch.get("startRow") or progress["nextStartRow"])
    end_row = args.end_row or int(last_batch.get("endRow") or progress["completedThroughRow"])

    workbook = openpyxl.load_workbook(workbook_path, data_only=True, read_only=True)
    sheet_names = ordered_range_sheets(workbook.sheetnames)
    if sheet not in sheet_names:
        raise SystemExit(f"Sheet not found: {sheet}")

    sheet_window = sheet_names[sheet_names.index(sheet) : sheet_names.index(sheet) + 2]
    words_by_sheet = {sheet_name: parse_sheet(workbook[sheet_name], sheet_name, max_row=(end_row + 5 if sheet_name == sheet else 10)) for sheet_name in sheet_window}
    batch_words = [
        word
        for word in words_by_sheet[sheet]
        if word.end_row >= start_row and word.start_row <= end_row
    ]

    words_doc = read_json(project_root / WORDS_SOURCE)
    roots_doc = read_json(project_root / ROOTS_SOURCE)
    words_by_slug = {word["slug"]: word for word in words_doc.get("words", [])}
    root_ids = {root["id"] for root in roots_doc.get("roots", [])}
    issues: list[str] = []
    warnings: list[str] = []

    check_duplicates(words_doc.get("words", []), "slug", "word slug", issues)
    check_duplicates(roots_doc.get("roots", []), "id", "root id", issues)

    expected_next = find_next_start(words_by_sheet, sheet_names, sheet, end_row)
    check_progress(progress, last_batch, sheet, start_row, end_row, expected_next, len(batch_words), issues, warnings)
    check_batch_words(project_root, batch_words, words_by_slug, root_ids, issues, warnings)

    result = {
        "batch": {
            "sheet": sheet,
            "startRow": start_row,
            "endRow": end_row,
            "excelRowsCovered": count_rows(batch_words),
            "logicalWords": len(batch_words),
            "firstWord": batch_words[0].term if batch_words else None,
            "lastWord": batch_words[-1].term if batch_words else None,
        },
        "nextStart": expected_next,
        "wordCount": len(words_doc.get("words", [])),
        "rootCount": len(roots_doc.get("roots", [])),
        "checkedSlugs": [word.slug for word in batch_words],
        "issues": issues,
        "warnings": warnings,
    }

    if args.json:
        print(json.dumps(result, ensure_ascii=False, indent=2))
    else:
        print_human_check(result)
    return 1 if issues else 0


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--project-root", default=".", help="Project root. Defaults to current directory.")
    parser.add_argument("--progress", default=str(DEFAULT_PROGRESS), help="Progress JSON path.")
    parser.add_argument("--workbook", help="Workbook path. Defaults to progress.workbook.")
    parser.add_argument("--sheet", help="Override checked sheet. Defaults to progress.lastBatch.sheet.")
    parser.add_argument("--start-row", type=int, help="Override checked start row.")
    parser.add_argument("--end-row", type=int, help="Override checked end row.")
    parser.add_argument("--json", action="store_true", help="Print machine-readable JSON.")
    return parser.parse_args()


def read_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def resolve_project_path(project_root: Path, path: Path) -> Path:
    return path if path.is_absolute() else project_root / path


def ordered_range_sheets(sheet_names: list[str]) -> list[str]:
    def key(name: str) -> tuple[int, str]:
        match = re.match(r"^(\d+)-\d+$", name)
        return (int(match.group(1)) if match else 10**9, name)

    return sorted(sheet_names, key=key)


def parse_sheet(worksheet: Any, sheet_name: str, max_row: int | None = None) -> list[LogicalWord]:
    headers = [cell.value for cell in worksheet[1]]
    missing = [header for header in REQUIRED_HEADERS if header not in headers]
    if missing:
        raise ValueError(f"{sheet_name} missing headers: {', '.join(missing)}")

    index = {header: headers.index(header) + 1 for header in REQUIRED_HEADERS}
    logical_words: list[LogicalWord] = []
    current: LogicalWord | None = None

    last_row = max_row if max_row is not None else worksheet.max_row
    for row_number in range(2, last_row + 1):
        raw_no = worksheet.cell(row_number, index["No."]).value
        raw_word = worksheet.cell(row_number, index["Word"]).value
        raw_pos = worksheet.cell(row_number, index["POS"]).value
        raw_meaning = worksheet.cell(row_number, index["Meaning"]).value
        raw_source = worksheet.cell(row_number, index["Source_Page"]).value

        if all(value is None for value in (raw_no, raw_word, raw_pos, raw_meaning, raw_source)):
            continue

        term = clean(raw_word)
        if term:
            current = LogicalWord(sheet_name, row_number, row_number, to_int(raw_no), term)
            logical_words.append(current)
        elif current is None:
            raise ValueError(f"{sheet_name} row {row_number} is a continuation row without a word")

        assert current is not None
        current.end_row = row_number
        current.rows.append(row_number)

    return logical_words


def to_int(value: Any) -> int | None:
    if value is None or value == "":
        return None
    return int(value)


def clean(value: Any) -> str:
    return "" if value is None else str(value).strip()


def slugify(term: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", term.lower()).strip("-")


def check_duplicates(entries: list[dict[str, Any]], key: str, label: str, issues: list[str]) -> None:
    seen: set[str] = set()
    for entry in entries:
        value = entry.get(key)
        if not value:
            continue
        if value in seen:
            issues.append(f"Duplicate {label}: {value}")
        seen.add(value)


def find_next_start(
    words_by_sheet: dict[str, list[LogicalWord]],
    sheet_names: list[str],
    end_sheet: str,
    end_row: int,
) -> dict[str, int | str | None]:
    for sheet_name in sheet_names[sheet_names.index(end_sheet) :]:
        for word in words_by_sheet.get(sheet_name, []):
            if sheet_name == end_sheet and word.start_row <= end_row:
                continue
            return {"sheet": sheet_name, "row": word.start_row, "no": word.number}
    return {"sheet": None, "row": None, "no": None}


def check_progress(
    progress: dict[str, Any],
    last_batch: dict[str, Any],
    sheet: str,
    start_row: int,
    end_row: int,
    expected_next: dict[str, int | str | None],
    logical_word_count: int,
    issues: list[str],
    warnings: list[str],
) -> None:
    expected_rows = last_batch.get("excelRowsCovered")
    if expected_rows is not None and int(expected_rows) != int(last_batch.get("endRow", end_row)) - int(last_batch.get("startRow", start_row)) + 1:
        warnings.append("lastBatch.excelRowsCovered differs from the inclusive row span; continuation blanks may need review.")

    if progress.get("completedThroughSheet") != sheet:
        issues.append(f"completedThroughSheet is {progress.get('completedThroughSheet')}, expected {sheet}")
    if int(progress.get("completedThroughRow", -1)) != end_row:
        issues.append(f"completedThroughRow is {progress.get('completedThroughRow')}, expected {end_row}")
    if progress.get("nextStartSheet") != expected_next["sheet"]:
        issues.append(f"nextStartSheet is {progress.get('nextStartSheet')}, expected {expected_next['sheet']}")
    if progress.get("nextStartRow") != expected_next["row"]:
        issues.append(f"nextStartRow is {progress.get('nextStartRow')}, expected {expected_next['row']}")
    if expected_next["no"] is not None and progress.get("nextStartNo") != expected_next["no"]:
        issues.append(f"nextStartNo is {progress.get('nextStartNo')}, expected {expected_next['no']}")
    if last_batch.get("sheet") and last_batch["sheet"] != sheet:
        issues.append(f"lastBatch.sheet is {last_batch['sheet']}, expected {sheet}")
    if last_batch.get("startRow") and int(last_batch["startRow"]) != start_row:
        issues.append(f"lastBatch.startRow is {last_batch['startRow']}, expected {start_row}")
    if last_batch.get("endRow") and int(last_batch["endRow"]) != end_row:
        issues.append(f"lastBatch.endRow is {last_batch['endRow']}, expected {end_row}")
    if last_batch.get("logicalWordsProcessed") and int(last_batch["logicalWordsProcessed"]) != logical_word_count:
        issues.append(
            f"lastBatch.logicalWordsProcessed is {last_batch['logicalWordsProcessed']}, "
            f"expected {logical_word_count}"
        )


def check_batch_words(
    project_root: Path,
    batch_words: list[LogicalWord],
    words_by_slug: dict[str, dict[str, Any]],
    root_ids: set[str],
    issues: list[str],
    warnings: list[str],
) -> None:
    if not batch_words:
        issues.append("No workbook words found for the requested batch range.")
        return

    for source_word in batch_words:
        entry = words_by_slug.get(source_word.slug)
        if entry is None:
            issues.append(f"Missing canonical word for workbook term: {source_word.term} ({source_word.slug})")
            continue

        generated_word = project_root / WORDS_OUTPUT / f"{source_word.slug}.yaml"
        if not generated_word.exists():
            issues.append(f"Missing generated word YAML: {generated_word}")

        for root_id in entry.get("rootIds", []):
            if root_id not in root_ids:
                issues.append(f"{source_word.slug} references missing rootId {root_id}")
            morpheme_file = project_root / MORPHEMES_OUTPUT / f"{safe_yaml_basename(strip_root_prefix(root_id))}.yaml"
            if not morpheme_file.exists():
                warnings.append(f"Missing generated morpheme YAML for {source_word.slug}: {morpheme_file}")

        for morpheme in entry.get("morphemes", []):
            root_id = morpheme.get("rootId")
            if root_id not in root_ids:
                issues.append(f"{source_word.slug}.{morpheme.get('form', 'morpheme')} references missing {root_id}")


def strip_root_prefix(root_id: str) -> str:
    return root_id.removeprefix("root:")



def safe_yaml_basename(root_id: str) -> str:
    reserved_windows_names = {
        "con",
        "prn",
        "aux",
        "nul",
        "com1",
        "com2",
        "com3",
        "com4",
        "com5",
        "com6",
        "com7",
        "com8",
        "com9",
        "lpt1",
        "lpt2",
        "lpt3",
        "lpt4",
        "lpt5",
        "lpt6",
        "lpt7",
        "lpt8",
        "lpt9",
    }
    return f"{root_id}-root" if root_id.lower() in reserved_windows_names else root_id

def count_rows(words: list[LogicalWord]) -> int:
    return sum(len(word.rows) for word in words)


def print_human_check(result: dict[str, Any]) -> None:
    batch = result["batch"]
    next_start = result["nextStart"]
    print("TOEFL 5400 batch check")
    print(
        f"Batch: {batch['sheet']} rows {batch['startRow']}-{batch['endRow']} "
        f"({batch['excelRowsCovered']} source rows, {batch['logicalWords']} logical words)"
    )
    print(f"First/last: {batch['firstWord']} -> {batch['lastWord']}")
    print(f"Next start: {next_start['sheet']} row {next_start['row']} / No. {next_start['no']}")
    print(f"Canonical totals: {result['wordCount']} words, {result['rootCount']} roots")
    print(f"Checked slugs: {', '.join(result['checkedSlugs'])}")
    if result["issues"]:
        print("")
        print("Issues:")
        for issue in result["issues"]:
            print(f"- {issue}")
    if result["warnings"]:
        print("")
        print("Warnings:")
        for warning in result["warnings"]:
            print(f"- {warning}")
    if not result["issues"]:
        print("")
        print("OK: batch source, canonical words, generated word files, root ids, and progress continuity passed.")


if __name__ == "__main__":
    raise SystemExit(main())




